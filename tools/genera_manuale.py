"""
Genera il manuale utente di DentalCare Pro dalle lezioni dell'Academy (il manuale ufficiale dentro il programma).

    python tools/genera_manuale.py <percorso del repo applicativo dentalcare>

Scrive in docs/05-Manuale-Utente/:
  - MANUALE_UTENTE.md          manuale integrale: tutti i moduli e le lezioni, nell'ordine dell'Academy
  - 01-Guida-Dottoressa.md     percorso del medico e dell'igienista con le lezioni che richiama
  - 02-Guida-Segretaria.md     percorso della segreteria con le lezioni che richiama
  - 03-Guida-Amministratore.md percorso del titolare con le lezioni che richiama
  - Manuale-Completo.md        sintesi per moduli, con le schermate
  - Manuale_Utente_DentalCare_Pro.docx  versione Word del manuale integrale
  - img/academy/*.png          schermate copiate dall'Academy

Le lezioni restano la sola fonte: per correggere il manuale si corregge la lezione e si rigenera.
"""
import io
import json
import re
import shutil
import sys
from datetime import date
from pathlib import Path

APP = Path(sys.argv[1] if len(sys.argv) > 1 else 'D:/dentalcare')
ACADEMY = APP / 'backend/src/main/resources/academy/v1'
IMAGES = APP / 'frontend/public/academy'
OUT = Path(__file__).resolve().parent.parent / 'docs/05-Manuale-Utente'
IMG_OUT = OUT / 'img/academy'

MESI = ['gennaio', 'febbraio', 'marzo', 'aprile', 'maggio', 'giugno', 'luglio', 'agosto', 'settembre',
        'ottobre', 'novembre', 'dicembre']
OGGI = date.today()
DATA = f'{OGGI.day} {MESI[OGGI.month - 1]} {OGGI.year}'

catalog = json.loads((ACADEMY / 'catalog.json').read_text(encoding='utf-8'))
modules = sorted(catalog['modules'], key=lambda m: m.get('sortOrder', 0))
for m in modules:
    m['lessons'] = sorted(m['lessons'], key=lambda l: l.get('sortOrder', 0))
lessons = {l['slug']: l for m in modules for l in m['lessons']}
module_of = {l['slug']: m for m in modules for l in m['lessons']}
number = {}
for i, m in enumerate(modules, 1):
    for j, l in enumerate(m['lessons'], 1):
        number[l['slug']] = f'{i}.{j}'


def body(slug):
    return (ACADEMY / 'lessons' / f'{slug}.md').read_text(encoding='utf-8').replace('\r\n', '\n').strip()


def links_in(text):
    return list(dict.fromkeys(re.findall(r'\]\(/academy/lezione/([a-z0-9-]+)\)', text)))


def convert(text, local, target_doc):
    """Link alle lezioni → ancore (nello stesso documento) o al manuale integrale; titoli abbassati di un livello."""
    def link(m):
        slug = m.group(1)
        if slug in local:
            return f'](#lezione-{slug})'
        return f']({target_doc}#lezione-{slug})'
    text = re.sub(r'\]\(/academy/lezione/([a-z0-9-]+)\)', link, text)
    text = re.sub(r'\]\(/academy/([^)]+\.png)\)', r'](img/academy/\1)', text)
    text = re.sub(r'^### Prova tu$', '#### Prova nel programma', text, flags=re.M)
    text = re.sub(r'^### ', '#### ', text, flags=re.M)
    return text


def lesson_block(slug, local, target_doc, heading='###'):
    l = lessons[slug]
    parts = [f'<a id="lezione-{slug}"></a>', '', f'{heading} {number[slug]} {l["title"]}', '',
             f'*{l["summary"]}*', '']
    if l.get('imageUrl'):
        name = l['imageUrl'].rsplit('/', 1)[1]
        caption = l.get('imageCaption') or l['title']
        parts += [f'![{caption}](img/academy/{name})', '', f'*{caption}*', '']
    parts += [convert(body(slug), local, target_doc), '']
    return '\n'.join(parts)


HEADER_NOTE = (f'> Versione {DATA}. Generato dalle lezioni dell\'**Academy**, il manuale ufficiale dentro il programma '
               '(`tools/genera_manuale.py`): per correggerlo si corregge la lezione e si rigenera. Le schermate vengono '
               'dall\'ambiente dimostrativo: nomi e dati dei pazienti sono fittizi.')


def write(name, text):
    (OUT / name).write_text(text.rstrip() + '\n', encoding='utf-8')


# ── schermate ─────────────────────────────────────────────────────────────
IMG_OUT.mkdir(parents=True, exist_ok=True)
used = {l['imageUrl'].rsplit('/', 1)[1] for l in lessons.values() if l.get('imageUrl')}
for slug in lessons:
    used |= set(re.findall(r'\]\(/academy/([^)]+\.png)\)', body(slug)))
for name in sorted(used):
    src = IMAGES / name
    if src.exists():
        shutil.copyfile(src, IMG_OUT / name)

# ── manuale integrale ─────────────────────────────────────────────────────
all_slugs = set(lessons)
out = ['# DentalCare Pro — Manuale utente', '', HEADER_NOTE, '', '## Indice', '']
for i, m in enumerate(modules, 1):
    out.append(f'{i}. [{m["title"]}](#modulo-{m["slug"]}) — {m["summary"]}')
    for l in m['lessons']:
        out.append(f'   - [{number[l["slug"]]} {l["title"]}](#lezione-{l["slug"]})')
out += ['', '---', '']
for i, m in enumerate(modules, 1):
    out += [f'<a id="modulo-{m["slug"]}"></a>', '', f'## {i}. {m["title"]}', '', f'*{m["summary"]}*', '']
    for l in m['lessons']:
        out.append(lesson_block(l['slug'], all_slugs, 'MANUALE_UTENTE.md'))
    out += ['---', '']
write('MANUALE_UTENTE.md', '\n'.join(out))
manual_text = '\n'.join(out)

# ── guide per ruolo ───────────────────────────────────────────────────────
GUIDES = [
    ('01-Guida-Dottoressa.md', 'Guida per il medico e l\'igienista', ['percorso-medico', 'percorso-igienista']),
    ('02-Guida-Segretaria.md', 'Guida per la segreteria', ['percorso-segreteria']),
    ('03-Guida-Amministratore.md', 'Guida per il titolare e l\'amministratore', ['percorso-titolare']),
]
for name, title, paths in GUIDES:
    order = list(paths)
    for p in paths:
        for s in links_in(body(p)):
            if s in lessons and s not in order:
                order.append(s)
    local = set(order)
    g = [f'# DentalCare Pro — {title}', '', HEADER_NOTE, '',
         'Il percorso dice cosa fare e in che ordine; le lezioni che richiama seguono, nell\'ordine in cui compaiono. '
         'Il resto è nel [manuale integrale](MANUALE_UTENTE.md).', '', '## Indice', '']
    for s in order:
        g.append(f'- [{lessons[s]["title"]}](#lezione-{s})')
    g += ['', '---', '']
    for s in order:
        g.append(lesson_block(s, local, 'MANUALE_UTENTE.md', heading='##'))
    write(name, '\n'.join(g))

# ── sintesi con le schermate ──────────────────────────────────────────────
c = ['# DentalCare Pro — Sintesi del manuale', '', HEADER_NOTE, '',
     'Ogni modulo con le sue lezioni in una riga e le schermate principali. Il testo completo è nel '
     '[manuale integrale](MANUALE_UTENTE.md).', '']
for i, m in enumerate(modules, 1):
    c += [f'## {i}. {m["title"]}', '', f'*{m["summary"]}*', '']
    for l in m['lessons']:
        c.append(f'- **[{l["title"]}](MANUALE_UTENTE.md#lezione-{l["slug"]})** — {l["summary"]}')
    c.append('')
    for l in m['lessons']:
        if l.get('imageUrl'):
            n = l['imageUrl'].rsplit('/', 1)[1]
            cap = l.get('imageCaption') or l['title']
            c += [f'![{cap}](img/academy/{n})', '', f'*{cap}*', '']
write('Manuale-Completo.md', '\n'.join(c))

# ── versione Word ─────────────────────────────────────────────────────────
from docx import Document  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)
INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def add_inline(par, text):
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**'):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith('`') and piece.endswith('`'):
            r = par.add_run(piece[1:-1])
            r.font.name = 'Consolas'
        elif piece.startswith('[') and '](' in piece:
            par.add_run(piece[1:piece.index('](')])
        elif piece.startswith('*') and piece.endswith('*') and len(piece) > 2:
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


lines = manual_text.split('\n')
i = 0
while i < len(lines):
    line = lines[i]
    s = line.strip()
    if not s or s.startswith('<a id=') or s == '---':
        i += 1
        continue
    if s.startswith('#'):
        level = len(s) - len(s.lstrip('#'))
        text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', s[level:].strip())
        doc.add_heading(text, level=min(level, 4) if level > 1 else 0)
    elif s.startswith('!['):
        m = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', s)
        path = OUT / m.group(2)
        if path.exists():
            doc.add_picture(str(path), width=Cm(16))
    elif s.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            cells = [x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{2,}:?', x) for x in cells):
                rows.append(cells)
            i += 1
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = 'Light Grid Accent 1'
        for r, cells in enumerate(rows):
            for k, cell in enumerate(cells):
                add_inline(table.cell(r, k).paragraphs[0], cell)
        continue
    elif re.match(r'^(\s*)[-*] ', line):
        indent = len(line) - len(line.lstrip())
        add_inline(doc.add_paragraph(style='List Bullet 2' if indent >= 2 else 'List Bullet'), re.sub(r'^\s*[-*] ', '', line))
    elif re.match(r'^\s*\d+\. ', line):
        add_inline(doc.add_paragraph(style='List Number'), re.sub(r'^\s*\d+\. ', '', line))
    elif s.startswith('>'):
        p = doc.add_paragraph(style='Intense Quote')
        add_inline(p, s.lstrip('> '))
    else:
        add_inline(doc.add_paragraph(), s)
    i += 1
doc.save(str(OUT / 'Manuale_Utente_DentalCare_Pro.docx'))

print(f'{len(modules)} moduli, {len(lessons)} lezioni, {len(used)} schermate -> {OUT}')
